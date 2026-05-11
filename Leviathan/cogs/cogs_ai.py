import discord
from discord.ext import commands
from groq import Groq
import os, datetime, io, json, re, asyncio, base64, random
import urllib.parse
import aiohttp

try:
    from docx import Document
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import openpyxl
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

OWNER_ID = 1136231768534569090
MODEL    = "llama-3.3-70b-versatile"
VISION_MODEL = "llama-3.2-90b-vision-preview"

PERSONALITY_FILE = "leviathan_personality.json"
ACCESS_FILE      = "ai_access.json"

DEFAULT_SYSTEM_PROMPT = (
    "You are Leviathan, an Ancient Dragon God and loyal familiar to Noelle Silva. "
    "You originate from the underwater Kingdom near the Cover Kingdom (Black Clover). "
    "Your personality: intelligent, sharp, slightly sassy, professional, and confident. "
    "You give concise, direct answers — short questions get short answers. "
    "You stay in character but do NOT pepper every message with ocean or sea metaphors. "
    "You speak plainly and helpfully, with a subtle edge of power. "
    "Never claim to be an AI model or any known product. "
    "When asked to generate files or images, describe what you are doing briefly, then produce it. "
    "Your knowledge has a cutoff of early 2024. If asked about recent events AND no web search "
    "results are provided in the message, say you may be outdated. However, if web search results "
    "ARE included in the message context, you MUST use them to answer — never say you lack internet "
    "access or are outdated when search results are present. Always answer from the provided results."
)

BOT_COMMANDS_REFERENCE = """
[LEVIATHAN BOT COMMANDS — know these and guide users to the right command when relevant]
WELCOME:       /welcome setup, /welcome toggle (enable/disable), /welcome status, /welcome preview
MODERATION:    /warn, /warnings, /clearwarnings, /kick, /ban, /timeout, /lock, /unlock, /slowmode, /clear
EMBED SYSTEM:  /embed, /editembed, /embedannounce, /embedrules, /embedwelcome, /embeddraft, /embedpost, /embedvars
ROLES:         /roleadd, /roleremove (or similar role commands)
UTILITY:       /serverinfo, /userinfo, /avatar, /uptime, /poll, /coinflip, /roll, /8ball
AI CONTROLS:   /aiallow, /airevoke, /aiallowlist, /aipersonality, /aireset, /aistatus
When a user asks you to perform a bot action (e.g. "turn off welcome", "disable welcome messages"),
tell them the exact command to use AND attempt to handle it via natural language if you can.
"""

# ── persistence ──────────────────────────────────────────────────────────────

def load_personality() -> str:
    if os.path.exists(PERSONALITY_FILE):
        try:
            with open(PERSONALITY_FILE) as f:
                return json.load(f).get("system_prompt", DEFAULT_SYSTEM_PROMPT)
        except Exception:
            pass
    return DEFAULT_SYSTEM_PROMPT

def save_personality(p: str):
    with open(PERSONALITY_FILE, "w") as f:
        json.dump({"system_prompt": p}, f, indent=2)

def load_access() -> set[int]:
    if os.path.exists(ACCESS_FILE):
        try:
            with open(ACCESS_FILE) as f:
                return set(json.load(f).get("allowed", []))
        except Exception:
            pass
    return set()

def save_access(allowed: set[int]):
    with open(ACCESS_FILE, "w") as f:
        json.dump({"allowed": list(allowed)}, f, indent=2)

# ── permission checks ─────────────────────────────────────────────────────────

def owner_only():
    async def predicate(interaction: discord.Interaction) -> bool:
        return interaction.user.id == OWNER_ID
    return discord.app_commands.check(predicate)

# ── helpers ───────────────────────────────────────────────────────────────────

def _dice_roll(expr: str) -> tuple[int, str]:
    """Parse XdY notation. Returns (total, detail_str)."""
    m = re.search(r"(\d*)d(\d+)", expr, re.I)
    if m:
        n  = int(m.group(1)) if m.group(1) else 1
        d  = int(m.group(2))
        n  = min(n, 20)
        rolls = [random.randint(1, d) for _ in range(n)]
        return sum(rolls), f"{n}d{d} → {rolls}"
    sides = int(re.search(r"\d+", expr).group()) if re.search(r"\d+", expr) else 6
    r = random.randint(1, sides)
    return r, f"d{sides} → {r}"

EIGHT_BALL = [
    "It is certain.", "It is decidedly so.", "Without a doubt.",
    "Yes, definitely.", "You may rely on it.", "As I see it, yes.",
    "Most likely.", "Outlook good.", "Yes.", "Signs point to yes.",
    "Reply hazy, try again.", "Ask again later.", "Better not tell you now.",
    "Cannot predict now.", "Concentrate and ask again.",
    "Don't count on it.", "My reply is no.", "My sources say no.",
    "Outlook not so good.", "Very doubtful.",
]

_SEARCH_TRIGGERS = re.compile(
    r"\b(latest|recent|news|today|current|now|2024|2025|2026|just|happened|"
    r"update|release|season|episode|chapter|announced|who\s+won|what\s+happened|"
    r"scores?|results?|standings?|trailer|dropped|premiere|streaming|"
    r"new\s+chapter|new\s+episode|new\s+season|right\s+now|this\s+week|"
    r"this\s+month|this\s+year|match|versus|vs\.?|beat|beaten|won|lost|"
    r"draw|tied|champion|winner|tournament|league|cup|final|fixture|"
    r"transfer|signing|roster|lineup)\b",
    re.I,
)

GEN_KEYWORDS = [
    "generate image", "create image", "draw", "make an image",
    "conjure an image", "paint", "generate a picture",
    "create a picture", "make a picture", "generate art",
]

MOD_PATTERNS = {
    "ban":     r"\bban\b",
    "kick":    r"\bkick\b",
    "timeout": r"\btimeout\b",
    "warn":    r"\bwarn\b",
    "lock":    r"\block\b",
    "unlock":  r"\bunlock\b",
}


class AICog(commands.Cog):
    def __init__(self, bot):
        self.bot            = bot
        self.client         = Groq(api_key=os.environ.get("GROQ_API_KEY", ""))
        self.system_prompt  = load_personality()
        self.history: dict[int, list[dict]] = {}
        self._pending:  dict[int, dict] = {}
        self.allowed_users: set[int] = load_access()

    # ── access ────────────────────────────────────────────────────────────────

    def _can_use(self, user: discord.User | discord.Member) -> bool:
        return user.id == OWNER_ID or user.id in self.allowed_users

    def _is_owner(self, user: discord.User | discord.Member) -> bool:
        return user.id == OWNER_ID

    def _should_respond(self, message: discord.Message) -> bool:
        if message.author.bot:
            return False
        if not self._can_use(message.author):
            return False
        if self.bot.user and self.bot.user in message.mentions:
            return True
        if isinstance(message.channel, discord.DMChannel):
            return True
        cl = message.content.lower()
        if "leviathan" in cl:
            return True
        if self._is_owner(message.author) and "levi" in cl:
            return True
        return False

    # ── AI call ───────────────────────────────────────────────────────────────

    def _runtime_context(self) -> str:
        utc = datetime.datetime.now(datetime.timezone.utc)

        def tz(offset_h: float, label: str) -> str:
            td = datetime.timedelta(hours=offset_h)
            t = datetime.datetime.now(datetime.timezone(td))
            return f"{label}: {t.strftime('%H:%M')}"

        zones = [
            (-10,   "Honolulu (HST)      "),
            (-8,    "Los Angeles (PST)   "),
            (-7,    "Denver (MST)        "),
            (-6,    "Chicago (CST)       "),
            (-5,    "New York (EST)      "),
            (-4,    "Halifax (AST)       "),
            (-3,    "São Paulo (BRT)     "),
            (0,     "London (GMT)        "),
            (1,     "Lagos/Paris (CET)   "),
            (2,     "Cairo/Joburg (EET)  "),
            (3,     "Moscow/Riyadh (MSK) "),
            (3.5,   "Tehran (IRST)       "),
            (4,     "Dubai (GST)         "),
            (5,     "Karachi (PKT)       "),
            (5.5,   "India (IST)         "),
            (5.75,  "Kathmandu (NPT)     "),
            (6,     "Dhaka (BST)         "),
            (7,     "Bangkok/Jakarta (ICT)"),
            (8,     "Singapore/Beijing   "),
            (9,     "Tokyo/Seoul (JST)   "),
            (9.5,   "Adelaide (ACST)     "),
            (10,    "Sydney (AEST)       "),
            (12,    "Auckland (NZST)     "),
        ]

        zone_lines = "\n".join(tz(off, lbl) for off, lbl in zones)
        return (
            f"\n[LIVE CONTEXT — injected at call time]\n"
            f"Current UTC: {utc.strftime('%A, %d %B %Y — %H:%M:%S UTC')}\n"
            f"World times right now:\n{zone_lines}\n"
            f"Use the exact times above when anyone asks what time it is anywhere."
            + BOT_COMMANDS_REFERENCE
        )

    async def _ai(self, user_id: int, prompt: str,
                  image_bytes: bytes | None = None,
                  mime: str = "image/jpeg") -> str:
        try:
            full_system = self.system_prompt + self._runtime_context()

            if image_bytes:
                b64 = base64.b64encode(image_bytes).decode()
                msgs = [
                    {"role": "system", "content": full_system},
                    {"role": "user", "content": [
                        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                        {"type": "text", "text": prompt},
                    ]},
                ]
                def _vc():
                    return self.client.chat.completions.create(
                        model=VISION_MODEL, messages=msgs, max_tokens=1024)
                r = await asyncio.to_thread(_vc)
                return r.choices[0].message.content or "…"

            hist  = self.history.get(user_id, [])
            msgs  = [{"role": "system", "content": full_system}]
            msgs += hist[-20:]
            msgs.append({"role": "user", "content": prompt})

            def _c():
                return self.client.chat.completions.create(
                    model=MODEL, messages=msgs, max_tokens=1024)
            r = await asyncio.to_thread(_c)
            reply = r.choices[0].message.content or "…"

            self.history.setdefault(user_id, [])
            self.history[user_id] += [
                {"role": "user",      "content": prompt},
                {"role": "assistant", "content": reply},
            ]
            self.history[user_id] = self.history[user_id][-40:]
            return reply
        except Exception as e:
            return f"Something went wrong: {e}"

    # ── web search ────────────────────────────────────────────────────────────

    async def _search(self, query: str) -> list[dict] | None:
        """Run a DuckDuckGo search. Returns raw result list or None."""
        try:
            def _do():
                from duckduckgo_search import DDGS
                return list(DDGS().text(query, max_results=5))
            results = await asyncio.to_thread(_do)
            return results if results else None
        except Exception:
            return None

    def _format_search_prompt(self, question: str, results: list[dict]) -> str:
        """
        Build a summarisation prompt. The AI is NEVER told to 'search' —
        it is only asked to summarise text we hand it. That always works.
        """
        snippets = []
        for r in results:
            title = r.get("title", "")
            body  = r.get("body", "")[:300]
            href  = r.get("href", "")
            snippets.append(f"Source: {title}\n{body}\n({href})")
        combined = "\n\n".join(snippets)
        return (
            f"Below are live web search results retrieved right now for the question:\n"
            f"'{question}'\n\n"
            f"--- SEARCH RESULTS START ---\n{combined}\n--- SEARCH RESULTS END ---\n\n"
            f"Using ONLY the information in the search results above, give a clear and "
            f"accurate answer to the question. Mention the source where relevant. "
            f"Do not say you cannot search — these results were already fetched for you."
        )

    # ── image generation ──────────────────────────────────────────────────────

    async def _gen_image(self, prompt: str) -> bytes | None:
        url = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(prompt)}?width=1024&height=1024&nologo=true"
        try:
            async with aiohttp.ClientSession() as s:
                async with s.get(url, timeout=aiohttp.ClientTimeout(total=60)) as r:
                    if r.status == 200:
                        return await r.read()
        except Exception:
            pass
        return None

    # ── file generation ───────────────────────────────────────────────────────

    async def _gen_docx(self, topic: str, channel) -> bool:
        if not DOCX_OK:
            await channel.send("❌ python-docx is not installed.")
            return False
        content = await self._ai(0, f"Write a well-structured document about: {topic}. Use clear headings and paragraphs. Plain text only.")
        buf = io.BytesIO()
        doc = Document()
        doc.add_heading(topic.title(), 0)
        for line in content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            else:
                doc.add_paragraph(line)
        doc.save(buf)
        buf.seek(0)
        fname = re.sub(r"[^\w\s-]", "", topic)[:40].strip().replace(" ", "_") + ".docx"
        await channel.send(f"📄 Here is your document on **{topic}**:", file=discord.File(buf, filename=fname))
        return True

    async def _gen_xlsx(self, description: str, channel) -> bool:
        if not XLSX_OK:
            await channel.send("❌ openpyxl is not installed.")
            return False
        raw = await self._ai(
            0,
            f"Create spreadsheet data for: {description}. "
            "Reply ONLY with CSV-style rows (comma-separated). First row = headers. "
            "Max 20 rows. No explanation, no markdown, just the data."
        )
        wb  = openpyxl.Workbook()
        ws  = wb.active
        for line in raw.strip().split("\n"):
            if line.strip():
                ws.append([c.strip().strip('"') for c in line.split(",")])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = re.sub(r"[^\w\s-]", "", description)[:40].strip().replace(" ", "_") + ".xlsx"
        await channel.send(f"📊 Here is your spreadsheet for **{description}**:", file=discord.File(buf, filename=fname))
        return True

    # ── moderation action handlers ────────────────────────────────────────────

    async def _request_mod(self, message: discord.Message, action: str, target: discord.Member | None) -> bool:
        if not isinstance(message.author, discord.Member):
            return False
        if not (message.author.guild_permissions.administrator or self._is_owner(message.author)):
            return False
        if not isinstance(message.channel, discord.TextChannel):
            return False

        target_str = target.display_name if target else "that member"
        phrase = (
            f"Shall I **{action}** {target_str}, master?"
            if self._is_owner(message.author)
            else f"Confirm: **{action}** {target_str}?"
        )
        self._pending[message.author.id] = {"action": action, "target": target, "channel": message.channel}
        await message.channel.send(f"{phrase} Reply `yes` to confirm or `no` to cancel.")
        return True

    async def _exec_mod(self, message: discord.Message) -> bool:
        pending = self._pending.pop(message.author.id, None)
        if not pending:
            return False
        if message.content.strip().lower() not in ("yes", "y", "confirm"):
            await message.channel.send("Action cancelled.")
            return True

        action = pending["action"]
        target = pending["target"]
        guild  = message.guild
        if not guild or not target:
            await message.channel.send("❌ Target not found.")
            return True

        try:
            if action == "ban":
                await target.ban(reason=f"AI request by {message.author}")
                await message.channel.send(f"✅ **{target.display_name}** has been banned.")
            elif action == "kick":
                await target.kick(reason=f"AI request by {message.author}")
                await message.channel.send(f"✅ **{target.display_name}** has been kicked.")
            elif action == "timeout":
                await target.timeout(discord.utils.utcnow() + datetime.timedelta(hours=1),
                                     reason=f"AI request by {message.author}")
                await message.channel.send(f"✅ **{target.display_name}** timed out for 1 hour.")
            elif action == "warn":
                from utils.warnings_manager import add_warning
                count = add_warning(guild.id, target.id, message.author.id, "Warned via AI")
                await message.channel.send(f"⚠️ **{target.display_name}** warned. Total: **{count}**.")
            elif action == "lock":
                ch = pending["channel"]
                if isinstance(ch, discord.TextChannel):
                    await ch.set_permissions(guild.default_role, send_messages=False)
                    await message.channel.send("🔒 Channel locked.")
            elif action == "unlock":
                ch = pending["channel"]
                if isinstance(ch, discord.TextChannel):
                    await ch.set_permissions(guild.default_role, send_messages=None)
                    await message.channel.send("🔓 Channel unlocked.")
        except discord.Forbidden:
            await message.channel.send("❌ I lack permission to do that.")
        except discord.HTTPException as e:
            await message.channel.send(f"❌ Error: {e}")
        return True

    # ── natural language command detection ────────────────────────────────────

    async def _handle_nl_command(self, message: discord.Message, text: str) -> bool:
        """Try to execute a natural language command. Returns True if handled."""
        tl = text.lower()
        guild = message.guild

        # ── roll dice ─────────────────────────────────────────────────────────
        if re.search(r"\broll\b", tl):
            expr = text
            total, detail = _dice_roll(expr)
            await message.channel.send(f"🎲 {detail} = **{total}**")
            return True

        # ── coin flip ─────────────────────────────────────────────────────────
        if re.search(r"\bcoin\b|\bflip\b|\bcoinflip\b", tl):
            result = random.choice(["**Heads** 🪙", "**Tails** 🪙"])
            await message.channel.send(f"The coin lands on {result}.")
            return True

        # ── 8ball ─────────────────────────────────────────────────────────────
        if re.search(r"\b8\s*ball\b|\beight\s*ball\b", tl):
            await message.channel.send(f"🎱 {random.choice(EIGHT_BALL)}")
            return True

        # ── avatar ────────────────────────────────────────────────────────────
        if re.search(r"\bavatar\b", tl) and message.mentions:
            u = message.mentions[0]
            e = discord.Embed(title=f"{u.display_name}'s Avatar", color=discord.Color.blurple())
            e.set_image(url=u.display_avatar.url)
            await message.channel.send(embed=e)
            return True

        # ── slowmode ──────────────────────────────────────────────────────────
        if re.search(r"\bslowmode\b", tl) and isinstance(message.channel, discord.TextChannel):
            nums = re.findall(r"\d+", text)
            secs = int(nums[0]) if nums else 5
            await message.channel.edit(slowmode_delay=secs)
            await message.channel.send(f"✅ Slowmode set to **{secs}s**.")
            return True

        # ── clear / purge ─────────────────────────────────────────────────────
        if re.search(r"\bclear\b|\bpurge\b|\bdelete\b.*\bmessage", tl):
            nums = re.findall(r"\d+", text)
            amt  = min(int(nums[0]), 100) if nums else 5
            if isinstance(message.channel, discord.TextChannel):
                deleted = await message.channel.purge(limit=amt + 1)
                m = await message.channel.send(f"✅ Deleted **{len(deleted)-1}** messages.")
                await asyncio.sleep(3)
                await m.delete()
            return True

        # ── poll ──────────────────────────────────────────────────────────────
        if re.search(r"\bpoll\b", tl):
            parts = re.split(r"\bpoll\b", text, maxsplit=1, flags=re.I)
            rest  = parts[1].strip() if len(parts) > 1 else ""
            segments = [s.strip() for s in re.split(r"\||\bor\b", rest) if s.strip()]
            question = segments[0] if segments else rest or "Poll"
            options  = segments[1:] if len(segments) > 1 else []
            e = discord.Embed(title=f"📊 {question}", color=discord.Color.blue())
            emojis = ["1️⃣","2️⃣","3️⃣","4️⃣","5️⃣","6️⃣","7️⃣","8️⃣","9️⃣","🔟"]
            if options:
                e.description = "\n".join(f"{emojis[i]} {o}" for i, o in enumerate(options[:10]))
            msg = await message.channel.send(embed=e)
            for i in range(min(len(options), 10) if options else 2):
                await msg.add_reaction(emojis[i] if options else ["👍","👎"][i])
            return True

        # ── embed in channel ──────────────────────────────────────────────────
        if re.search(r"\bembed\b", tl):
            target_ch = message.channel_mentions[0] if message.channel_mentions else message.channel
            title_m   = re.search(r'(?:titled?|called?|saying)\s+["\']?(.+?)["\']?(?:\s+in\s|$)', text, re.I)
            title     = title_m.group(1).strip() if title_m else "Embed"
            desc_m    = re.search(r'(?:with\s+(?:description|content|text)|description:?)\s+["\']?(.+?)["\']?(?:\s+in\s|$)', text, re.I)
            desc      = desc_m.group(1).strip() if desc_m else ""
            e = discord.Embed(title=title, description=desc, color=discord.Color.blurple())
            if isinstance(target_ch, discord.TextChannel):
                await target_ch.send(embed=e)
                if target_ch != message.channel:
                    await message.channel.send(f"✅ Embed posted in {target_ch.mention}.")
            return True

        # ── web search command ────────────────────────────────────────────────
        search_m = re.match(r"^search\s+(.+)", tl.strip(), re.I)
        if search_m or re.search(r"\bsearch\s+(for\s+|up\s+)?(.+)", tl, re.I):
            raw = re.sub(r"^search\s+(for\s+|up\s+)?", "", text.strip(), flags=re.I).strip()
            query = raw or text
            status_msg = await message.channel.send(f"🔍 Searching for: **{query}**…")
            results = await self._search(query)
            if not results:
                await status_msg.edit(content="❌ No results found for that query.")
                return True
            prompt = self._format_search_prompt(query, results)
            async with message.channel.typing():
                reply = await self._ai(message.author.id, prompt)
            await status_msg.delete()
            for chunk in [reply[i:i+1990] for i in range(0, len(reply), 1990)]:
                await message.channel.send(chunk)
            return True

        # ── create document ───────────────────────────────────────────────────
        if re.search(r"\bcreate\b.*\b(doc|document|word)\b|\b(doc|document)\b.*\bcreate\b", tl):
            topic_m = re.search(r"(?:about|on|for|regarding)\s+(.+)", text, re.I)
            topic   = topic_m.group(1).strip() if topic_m else text
            await message.channel.send("📄 Writing your document…")
            await self._gen_docx(topic, message.channel)
            return True

        # ── create spreadsheet ────────────────────────────────────────────────
        if re.search(r"\bcreate\b.*\b(spreadsheet|excel|xlsx|sheet)\b|\b(spreadsheet|excel)\b.*\bcreate\b", tl):
            desc_m = re.search(r"(?:with|about|for|on)\s+(.+)", text, re.I)
            desc   = desc_m.group(1).strip() if desc_m else text
            await message.channel.send("📊 Generating your spreadsheet…")
            await self._gen_xlsx(desc, message.channel)
            return True

        # ── image generation ──────────────────────────────────────────────────
        if any(kw in tl for kw in GEN_KEYWORDS):
            prompt = text
            for kw in sorted(GEN_KEYWORDS, key=len, reverse=True):
                prompt = re.sub(re.escape(kw), "", prompt, flags=re.I).strip()
            await message.channel.send("🎨 Conjuring your image…")
            data = await self._gen_image(prompt or text)
            if data:
                await message.channel.send(file=discord.File(io.BytesIO(data), "image.png"))
            else:
                await message.channel.send("❌ Could not generate the image.")
            return True

        # ── welcome toggle ────────────────────────────────────────────────────
        if re.search(r"\bwelcome\b", tl) and re.search(r"\b(toggle|enable|disable|turn\s+on|turn\s+off|off|on)\b", tl):
            welcome_cog = self.bot.cogs.get("WelcomeCog")
            if welcome_cog is None:
                await message.channel.send("❌ Welcome system not loaded.")
                return True
            guild = message.guild
            if guild is None:
                return True
            guild_id = str(guild.id)
            cfg = welcome_cog.config.get(guild_id)
            if not cfg:
                await message.channel.send("❌ Welcome system isn't configured yet. Use `/welcome setup` first.")
                return True
            # Explicit enable/disable intent
            want_enable  = bool(re.search(r"\b(enable|turn\s+on|\bon\b)\b", tl))
            want_disable = bool(re.search(r"\b(disable|turn\s+off|\boff\b)\b", tl))
            current = cfg.get("enabled", True)
            if want_enable:
                cfg["enabled"] = True
            elif want_disable:
                cfg["enabled"] = False
            else:
                cfg["enabled"] = not current  # plain "toggle"
            welcome_cog.config[guild_id] = cfg
            from cogs.cogs_welcome import save_config
            save_config(welcome_cog.config)
            state = "✅ enabled" if cfg["enabled"] else "🔴 disabled"
            await message.channel.send(f"Auto-welcome is now **{state}**.")
            return True

        # ── moderation ────────────────────────────────────────────────────────
        for action, pat in MOD_PATTERNS.items():
            if re.search(pat, tl):
                target = message.mentions[0] if message.mentions else None
                if target or action in ("lock", "unlock"):
                    return await self._request_mod(message, action, target)

        return False

    # ── main listener ─────────────────────────────────────────────────────────

    @commands.Cog.listener()
    async def on_message(self, message: discord.Message):
        if message.author.bot:
            return

        # pending mod confirmation
        if message.author.id in self._pending:
            await self._exec_mod(message)
            return

        if not self._should_respond(message):
            return

        # clean content
        content = message.content
        if self.bot.user:
            content = content.replace(f"<@{self.bot.user.id}>", "").replace(f"<@!{self.bot.user.id}>", "").strip()
        if not content and not message.attachments:
            return

        # try NL command first (owner only)
        if self._is_owner(message.author):
            async with message.channel.typing():
                handled = await self._handle_nl_command(message, content)
                if handled:
                    return

        # attach image if present
        image_bytes: bytes | None = None
        mime = "image/jpeg"
        for att in message.attachments:
            fn = att.filename.lower()
            if any(fn.endswith(x) for x in [".png", ".jpg", ".jpeg", ".webp", ".gif"]):
                try:
                    image_bytes = await att.read()
                    if fn.endswith(".png"):   mime = "image/png"
                    elif fn.endswith(".gif"): mime = "image/gif"
                    elif fn.endswith(".webp"): mime = "image/webp"
                    break
                except Exception:
                    pass

        async with message.channel.typing():
            # Auto-search for queries that look like they need live data
            prompt = content
            if not image_bytes and _SEARCH_TRIGGERS.search(content):
                results = await self._search(content)
                if results:
                    prompt = self._format_search_prompt(content, results)
            reply = await self._ai(message.author.id, prompt, image_bytes, mime)

        chunks = [reply[i:i+1990] for i in range(0, len(reply), 1990)]
        for chunk in chunks:
            await message.channel.send(chunk)

    # ── slash commands ────────────────────────────────────────────────────────

    @discord.app_commands.command(name="aiallow", description="Grant a user access to Leviathan's AI.")
    @owner_only()
    @discord.app_commands.describe(user="User to allow")
    async def aiallow(self, interaction: discord.Interaction, user: discord.Member):
        self.allowed_users.add(user.id)
        save_access(self.allowed_users)
        await interaction.response.send_message(f"✅ **{user.display_name}** can now use Leviathan's AI.", ephemeral=True)

    @discord.app_commands.command(name="airevoke", description="Remove a user's access to Leviathan's AI.")
    @owner_only()
    @discord.app_commands.describe(user="User to revoke")
    async def airevoke(self, interaction: discord.Interaction, user: discord.Member):
        self.allowed_users.discard(user.id)
        save_access(self.allowed_users)
        await interaction.response.send_message(f"✅ **{user.display_name}**'s AI access removed.", ephemeral=True)

    @discord.app_commands.command(name="aiallowlist", description="View all users with AI access.")
    @owner_only()
    async def aiallowlist(self, interaction: discord.Interaction):
        if not self.allowed_users:
            await interaction.response.send_message("No additional users have AI access (owner only).", ephemeral=True)
            return
        lines = []
        for uid in self.allowed_users:
            u = interaction.guild.get_member(uid) if interaction.guild else None
            lines.append(f"• {u.display_name} ({uid})" if u else f"• {uid}")
        await interaction.response.send_message("**AI Access List:**\n" + "\n".join(lines), ephemeral=True)

    @discord.app_commands.command(name="aipersonality", description="View or edit Leviathan's personality.")
    @owner_only()
    @discord.app_commands.describe(new_prompt="New system prompt (blank = view current)")
    async def aipersonality(self, interaction: discord.Interaction, new_prompt: str | None = None):
        if new_prompt is None:
            e = discord.Embed(title="🐉 Current Personality",
                              description=f"```{self.system_prompt[:3900]}```",
                              color=discord.Color.blue())
            await interaction.response.send_message(embed=e, ephemeral=True)
        else:
            self.system_prompt = new_prompt
            save_personality(new_prompt)
            self.history.clear()
            await interaction.response.send_message("✅ Personality updated.", ephemeral=True)

    @discord.app_commands.command(name="aireset", description="Clear Leviathan's conversation memory.")
    @owner_only()
    @discord.app_commands.describe(user="Clear for one user only (blank = clear all)")
    async def aireset(self, interaction: discord.Interaction, user: discord.Member | None = None):
        if user:
            self.history.pop(user.id, None)
            await interaction.response.send_message(f"✅ Memory cleared for **{user.display_name}**.", ephemeral=True)
        else:
            self.history.clear()
            await interaction.response.send_message("✅ All memory cleared.", ephemeral=True)

    @discord.app_commands.command(name="aistatus", description="Check Leviathan's AI status.")
    async def aistatus(self, interaction: discord.Interaction):
        if not self._can_use(interaction.user):
            await interaction.response.send_message("❌ You don't have access to Leviathan's AI.", ephemeral=True)
            return
        key = bool(os.environ.get("GROQ_API_KEY"))
        e = discord.Embed(title="🐉 Leviathan AI Status",
                          color=discord.Color.green() if key else discord.Color.red())
        e.add_field(name="Provider",   value="✅ Groq (Free)"      if key else "❌ No Key", inline=True)
        e.add_field(name="Model",      value=MODEL,                                         inline=True)
        e.add_field(name="Convos",     value=str(len(self.history)),                        inline=True)
        e.add_field(name="Images",     value="✅ Pollinations AI",                           inline=True)
        e.add_field(name="Daily limit",value="~14,400 req/day (Groq free tier)",            inline=False)
        await interaction.response.send_message(embed=e, ephemeral=True)

    async def cog_app_command_error(self, interaction: discord.Interaction, error: discord.app_commands.AppCommandError):
        msg = "❌ Owner only." if isinstance(error, discord.app_commands.CheckFailure) else f"❌ {error}"
        if interaction.response.is_done():
            await interaction.followup.send(msg, ephemeral=True)
        else:
            await interaction.response.send_message(msg, ephemeral=True)


async def setup(bot):
    await bot.add_cog(AICog(bot))
